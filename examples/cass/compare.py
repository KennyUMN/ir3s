"""Compare a hand-built CASS model with what ir3s reads from the PDFs.

For every figure in the '3 Statements' sheet of CASS_FINMOD.xlsx, look for
the same number in the audited report ir3s read for that year. A figure
matches when it equals one printed line, or the sum of a run of adjacent
printed lines (the model often merges "related parties" and "third parties").
A figure that only appears in the next year's report matches a restated
number. Each match names the PDF page and the printed line, so anyone can
open the report and check it.

The model's cash flow statement is derived from balance-sheet changes, not
copied from the report, so only its opening and closing cash are compared.

    .venv/bin/python examples/cass/compare.py [out.xlsx]
"""
import os
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent.parent))
from ir3s.extractor import extract_statements  # noqa: E402

MODEL = os.environ.get("CASS_MODEL", os.path.expanduser("~/Downloads/Documents/CASS_FINMOD.xlsx"))
PDF_DIR = os.environ.get("CASS_PDF_DIR", os.path.expanduser("~/.hermes/cache/scratch/audit/cass"))
OUT = sys.argv[1] if len(sys.argv) > 1 else str(HERE / "CASS_reconciliation.xlsx")
SHEET = "3 Statements"
YEARS = {"C": 2019, "D": 2020, "E": 2021, "F": 2022, "G": 2023, "H": 2024}
# (statement, first row, last row, ir3s kind); the check rows are left out.
PARTS = [
    ("Neraca", 8, 63, "balance_sheet"),
    ("Laba Rugi", 69, 86, "income_statement"),
    ("Arus Kas", 143, 144, "cash_flow"),
]
SKIP_ROWS = {85}  # "% of Net profit", a ratio
MAX_RUN = 4  # longest run of adjacent printed lines a model figure may merge

MATCH, RESTATED, MISSING = "cocok", "cocok (angka disajikan kembali)", "tidak ditemukan"

# Why each figure that was not found differs, checked by hand against the PDF.
# Keyed by (model row, year); a year of None covers every year.
SPLIT = ("Laporan hanya mencetak satu baris 'Beban usaha'. Model memecahnya jadi beban pokok "
         "dan beban usaha dari catatan (sheet 'Financial Modelling'), jadi angka ini tidak ada di muka laporan.")
EXPLAIN = {
    (70, None): SPLIT,
    (72, None): SPLIT,
    (71, None): "Laba bruto tidak dicetak di laporan; mengikuti pemecahan beban di atas.",
    (73, 2019): "Laporan: laba usaha 435.326. Pemecahan beban di model berjumlah 1.759.221, "
                "laporan mencetak beban usaha 1.758.980: selisih 241.",
    (78, 2019): "Laporan: 145.112. Selisih 241 yang sama dari laba usaha.",
    (73, 2022): "Laporan LK 2022: laba usaha 400.297. Model memakai pendapatan LK 2022 (1.737.746) tapi beban "
                "dari LK 2023 yang disajikan kembali (1.313.322), jadi laba usaha lebih tinggi 24.127.",
    (78, 2022): "Laporan LK 2022: 399.433. Selisih 24.127 yang sama dari laba usaha.",
    (80, 2022): "Laporan LK 2022: 289.798, dan baris 86 model sendiri juga 289.798. Selisih 24.127 dari laba usaha.",
    (77, 2023): "Model memasukkan rugi operasi yang dihentikan (3.052) ke beban operasi lain. Laporan "
                "mencetaknya di bawah pajak, jadi beban operasi lain 77.354.",
    (78, 2023): "Laporan: 565.962. Selisih 3.052 dari rugi operasi yang dihentikan; laba bersih tetap sama.",
    (76, 2024): "Model memasukkan laba operasi yang dihentikan (1.596) ke pendapatan lain. Laporan "
                "mencetaknya di bawah pajak, jadi pendapatan operasi lain 273.579.",
    (78, 2024): "Laporan: 901.508. Selisih 1.596 dari laba operasi yang dihentikan; laba bersih tetap sama.",
}


def load_reports():
    kinds = ("balance_sheet", "income_statement", "cash_flow")
    return {y: extract_statements(os.path.join(PDF_DIR, f"cass_fs_{y}.pdf"), kinds)
            for y in range(min(YEARS.values()), max(YEARS.values()) + 2)
            if os.path.exists(os.path.join(PDF_DIR, f"cass_fs_{y}.pdf"))}


def find(value, statement, year):
    """The printed lines whose figure, alone or summed in a run, equals value."""
    if statement is None:
        return None
    lines = [(r, r.value_for(year)) for r in statement.rows]
    lines = [(r, v) for r, v in lines if v is not None]
    for size in range(1, MAX_RUN + 1):
        for i in range(len(lines) - size + 1):
            run = lines[i:i + size]
            if abs(sum(v for _, v in run)) == abs(value):
                return [r for r, _ in run]
    return None


def compare():
    formulas = openpyxl.load_workbook(MODEL)[SHEET]
    values = openpyxl.load_workbook(MODEL, data_only=True)[SHEET]
    reports = load_reports()
    results = []
    for part, first, last, kind in PARTS:
        for row in range(first, last + 1):
            label = values[f"B{row}"].value
            if row in SKIP_ROWS or not label:
                continue
            for col, year in YEARS.items():
                value = values[f"{col}{row}"].value
                if not isinstance(value, (int, float)):
                    continue
                value = Decimal(str(value))
                status, source, lines = MISSING, None, None
                for report, found_status in ((year, MATCH), (year + 1, RESTATED)):
                    statement = reports.get(report, {}).get(kind)
                    lines = find(value, statement, year) if value else None
                    if lines:
                        status, source = found_status, report
                        break
                if not value:
                    status = "nol di model"
                results.append({
                    "part": part, "row": row, "label": label.strip(), "year": year,
                    "value": value, "formula": formulas[f"{col}{row}"].value,
                    "status": status, "source": source, "lines": lines or [],
                    "why": EXPLAIN.get((row, year)) or EXPLAIN.get((row, None)) if status == MISSING else None,
                })
    return results


NAVY = PatternFill("solid", fgColor="393E5C")
FILLS = {MATCH: "C6EFCE", RESTATED: "FFEB9C", MISSING: "FFC7CE", "nol di model": "EDEDED"}
FONT = "DM Sans"


def header(sh, cells):
    sh.append(cells)
    for cell in sh[sh.max_row]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = NAVY


def write(results):
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Ringkasan"
    summary.append(["Model CASS (diisi manual) dibandingkan dengan angka yang dibaca ir3s dari laporan keuangan auditan"])
    summary["A1"].font = Font(name=FONT, bold=True, size=13)
    summary.append([])
    header(summary, ["Laporan", "Tahun", "Angka di model", "Cocok", "Cocok (disajikan kembali)", "Tidak ditemukan", "% cocok"])
    for part, *_ in PARTS:
        for year in YEARS.values():
            rows = [r for r in results if r["part"] == part and r["year"] == year and r["value"]]
            n = {s: sum(r["status"] == s for r in rows) for s in (MATCH, RESTATED, MISSING)}
            summary.append([part, year, len(rows), n[MATCH], n[RESTATED], n[MISSING],
                            (n[MATCH] + n[RESTATED]) / len(rows) if rows else None])
            summary.cell(row=summary.max_row, column=7).number_format = "0%"
    for row in summary.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name=FONT)
    summary.append([])
    for line in [
        "Cara membaca: setiap angka di sheet '3 Statements' dicari di laporan keuangan auditan tahun yang sama.",
        "'Cocok' berarti angka itu tercetak di laporan, sendiri atau sebagai jumlah beberapa baris yang berdampingan.",
        "Sheet 'Detail' menyebut halaman PDF dan baris tercetaknya, jadi setiap angka bisa dicek langsung.",
        "Arus kas di model dihitung dari perubahan neraca, bukan disalin dari laporan; hanya kas awal dan akhir dibandingkan.",
        "Setiap angka yang tidak ditemukan sudah dicek manual ke PDF; alasannya ada di kolom 'Penjelasan' sheet 'Detail'.",
    ]:
        summary.append([line])
        summary.cell(row=summary.max_row, column=1).font = Font(name=FONT, size=10)
    summary.column_dimensions["A"].width = 14
    for c in "BCDEFG":
        summary.column_dimensions[c].width = 16

    detail = wb.create_sheet("Detail")
    header(detail, ["Laporan", "Baris model", "Pos", "Tahun", "Angka di model", "Rumus di model", "Hasil",
                    "Sumber (LK tahun)", "Halaman PDF", "Baris tercetak di laporan", "Angka tercetak", "Penjelasan"])
    for r in results:
        detail.append([
            r["part"], r["row"], r["label"], r["year"], r["value"],
            r["formula"] if isinstance(r["formula"], str) else None, r["status"],
            f"LK {r['source']}" if r["source"] else None,
            ", ".join(sorted({str(line.page) for line in r["lines"]})) or None,
            " + ".join(line.label[:70] for line in r["lines"]) or None,
            " + ".join(c.raw for line in r["lines"] for c in line.cells if c.year == r["year"]) or None,
            r["why"],
        ])
        for cell in detail[detail.max_row]:
            cell.font = Font(name=FONT, size=10)
        detail.cell(row=detail.max_row, column=5).number_format = '#,##0;(#,##0);"-"'
        detail.cell(row=detail.max_row, column=7).fill = PatternFill("solid", fgColor=FILLS[r["status"]])
    for col, width in zip("ABCDEFGHIJKL", (11, 8, 40, 7, 15, 28, 28, 12, 10, 70, 24, 90)):
        detail.column_dimensions[col].width = width
    for cell in detail["L"]:
        cell.alignment = Alignment(wrap_text=True)
    detail.freeze_panes = "E2"
    detail.auto_filter.ref = detail.dimensions
    for sh in (summary, detail):
        sh.sheet_view.showGridLines = False
        for row in sh.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    wb.save(OUT)


if __name__ == "__main__":
    results = compare()
    write(results)
    counted = [r for r in results if r["value"]]
    for status in (MATCH, RESTATED, MISSING):
        print(f"{status:32} {sum(r['status'] == status for r in counted)}")
    print("total figures compared:", len(counted), "| saved", OUT)
    for r in counted:
        if r["status"] == MISSING:
            print(f"  MISSING {r['part']:9} row {r['row']:3} {r['year']} {r['label'][:40]:40} {r['value']:>12} {r['formula']}")
