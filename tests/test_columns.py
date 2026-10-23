"""Which column a figure belongs to: the year and the period it covers.

Each case is a real statement that was read wrong: a figure landed under
another year, a comparative year was dropped, or a six- or nine-month
figure was read as a full year. The year comes from the header printed above
the figure's column, so the order the columns are printed in does not matter.
"""
from __future__ import annotations

import shutil
from decimal import Decimal
from pathlib import Path

import pymupdf
import pytest

from ir3s.extractor import extract_statements, read_page

SCRATCH = Path("/Users/kennyvws/.hermes/cache/scratch")
AUDIT = SCRATCH / "audit"
BCA = SCRATCH / "fsm"


def _statements(path: Path, kinds=("balance_sheet", "income_statement", "cash_flow")):
    if not path.exists():
        pytest.skip(f"{path.name} not present")
    return extract_statements(str(path), kinds)


def _row_with(statement, value: int):
    return next((row for row in statement.rows
                 if any(cell.value == value for cell in row.cells)), None)


def _cells(row) -> dict:
    return {(cell.year, cell.interim): cell.value for cell in row.cells}


@pytest.mark.parametrize("order", [(2019, 2018), (2018, 2019)])
def test_the_header_above_a_column_names_its_year_in_either_order(order):
    """2019 | 2018 and 2018 | 2019 both put each figure under its own year."""
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((60, 100), "Catatan/Notes")
    first, second = order
    page.insert_text((300, 120), str(first))
    page.insert_text((420, 120), str(second))
    figures = {2019: ("1.000", "2.500"), 2018: ("900", "2.000")}
    page.insert_text((60, 150), "Kas")
    page.insert_text((300, 150), figures[first][0])
    page.insert_text((420, 150), figures[second][0])
    page.insert_text((60, 170), "Jumlah aset")
    page.insert_text((300, 170), figures[first][1])
    page.insert_text((420, 170), figures[second][1])
    for y in range(190, 260, 20):
        page.insert_text((60, y), f"Akun {y}")
        page.insert_text((300, y), "10")
        page.insert_text((420, y), "20")

    years, rows = read_page(page)

    assert years == [2018, 2019]
    cash = rows[0]
    assert cash.value_for(2019) == 1000
    assert cash.value_for(2018) == 900


def test_a_stacked_header_does_not_shift_the_years():
    """CASS 2019 prints `2018 (restated)` four lines above `2019 | Notes`.

    The comparative was dropped and the 2019 figures were read as 2018.
    """
    income = _statements(AUDIT / "cass" / "cass_fs_2019.pdf")["income_statement"]
    revenue = _row_with(income, 2194306)
    assert revenue is not None
    assert revenue.value_for(2019) == 2194306
    assert revenue.value_for(2018) == 2200342


def test_an_opening_balance_column_is_the_prior_year_end():
    """CASS 2019's third column is 1 January 2018, which is 31 December 2017.

    It was read as 2018, and the real 31 December 2018 column was dropped.
    """
    balance = _statements(AUDIT / "cass" / "cass_fs_2019.pdf")["balance_sheet"]
    cash = _row_with(balance, 265015)
    assert _cells(cash) == {
        (2019, False): 265015, (2018, False): 200514, (2017, False): 136191,
    }


def test_a_notes_column_between_the_years_does_not_hide_the_comparative():
    """CASS 2022's cash flow prints 2022 | Notes | 2021.

    The 2021 figures were left inside the labels, and a row whose 2022 figure
    is a dash was glued onto the net investing total below it.
    """
    cash_flow = _statements(AUDIT / "cass" / "cass_fs_2022.pdf")["cash_flow"]
    receipts = _row_with(cash_flow, 1630002)
    assert _cells(receipts) == {(2022, False): 1630002, (2021, False): 1422880}
    assert "1.422.880" not in receipts.label
    investing = _row_with(cash_flow, -133947)
    assert investing.value_for(2021) == -6523
    assert "pihak berelasi" not in investing.label


def test_a_nine_month_column_is_not_read_as_the_year():
    """WBSA's Sep-25 report prints 9M 2025, 9M 2024, FY 2024, 2023, 2022.

    Both 2024 columns were tagged 2024, and the nine-month one came first.
    """
    income = _statements(AUDIT / "wbsa" / "wbsa_fs_sep2025.pdf")["income_statement"]
    assert income.years == [2022, 2023, 2024]
    revenue = _row_with(income, 1772043819)
    assert revenue.value_for(2024) == 1772043819
    assert _cells(revenue)[(2024, True)] == 1196883496
    assert _cells(revenue)[(2025, True)] == 1544766349
    assert revenue.value_for(2025) is None


def test_a_six_month_report_has_no_full_year_income():
    """WBSA's Q2 2026 report: the six months to June 2026 were read as 2025."""
    statements = _statements(AUDIT / "wbsa" / "wbsa_fs_q2_2026.pdf")
    income = statements["income_statement"]
    revenue = _row_with(income, 1948668952)
    assert _cells(revenue) == {(2026, True): 1948668952, (2025, True): 1366180924}
    assert income.years == []
    balance = statements["balance_sheet"]
    assets = _row_with(balance, 1717248401)
    assert _cells(assets) == {(2026, True): 1717248401, (2025, False): 1601613967}
    assert balance.years == [2025]


def test_an_interim_bank_report_keeps_the_full_years_apart():
    """Superbank's June 2025 report: the six-month figures are not 2025."""
    statements = _statements(AUDIT / "superbank" / "sbnr_audited_2.pdf")
    assert statements["balance_sheet"].years == [2022, 2023, 2024]
    receipts = _row_with(statements["cash_flow"], 846966)
    assert _cells(receipts) == {
        (2025, True): 846966, (2024, True): 297007,
        (2024, False): 766799, (2023, False): 280141, (2022, False): 146363,
    }


def test_a_group_member_after_a_wrapped_group_name_carries_the_name():
    """BCA 2024 prints Loans receivable over three lines, then its members.

    Third parties, the larger member, lost the group name.
    """
    balance = _statements(BCA / "bca_fs_2024.pdf", ("balance_sheet",))["balance_sheet"]
    third = _row_with(balance, 861511753)
    assert third.label.startswith("Loans receivable"), third.label
    related = _row_with(balance, 7174457)
    assert related.label.startswith("Loans receivable"), related.label
    assert related.label.endswith("Related parties"), related.label


def test_a_fill_leaves_the_years_it_did_not_read():
    """Filling one report must not erase the years another report filled.

    The writer used to empty every historical column of a statement's rows,
    so a 2019-2026 model could not be filled one report at a time.
    """
    import openpyxl

    from ir3s.sheet_writer import open_template, write_values

    source = SCRATCH / "fsm" / "template.xlsx"
    if not source.exists():
        pytest.skip("template not present")
    copy = Path(__file__).parent / "_years_test.xlsx"
    shutil.copyfile(source, copy)
    try:
        template = open_template(copy)
        sheet = template.sheet("3 Statements")
        columns = template.historical_columns(sheet, 5)
        kept_year, written_year = sorted(columns)[0], sorted(columns)[1]
        sheet.cell(row=10, column=columns[kept_year]).value = 123
        sheet.cell(row=10, column=columns[written_year]).value = 456
        write_values(
            template, "3 Statements", 5,
            {36: {written_year: Decimal(100)}},
            force=True, label_for_row={36: "row 36"},
            clear_rows=((6, 67, frozenset({written_year})),),
        )
        template.workbook.save(copy)
        sheet = openpyxl.load_workbook(copy)["3 Statements"]
        assert sheet.cell(row=10, column=columns[kept_year]).value == 123
        assert sheet.cell(row=10, column=columns[written_year]).value is None
        assert sheet.cell(row=36, column=columns[written_year]).value == 100
    finally:
        copy.unlink(missing_ok=True)


def test_a_bare_parenthesis_is_not_a_year():
    """DEWA 2025 prints a lone ")" in its header. It crashed the reader."""
    from ir3s.extractor import parse_year

    assert parse_year(")") is None
    assert parse_year("2024*)") == 2024


def test_a_header_with_both_dates_names_the_year_end():
    """DEWA 2025's third column is "1 January 2024/31 December 2023".

    It was read as 2022, a year the report does not print.
    """
    balance = _statements(AUDIT / "dewa" / "dewa_fs_2025.pdf")["balance_sheet"]
    assert balance.years == [2023, 2024, 2025]
    assert balance.unit == "thousands"
    cash = _row_with(balance, 1590630399)
    assert _cells(cash) == {
        (2025, False): 1590630399, (2024, False): 346530081, (2023, False): 386856009,
    }


def test_a_dollar_report_is_not_written_into_a_rupiah_template():
    """DEWA reported in US dollars until 2022 and in thousand rupiah after."""
    from ir3s.schemas import ModelSpec

    balance = _statements(AUDIT / "dewa" / "dewa_fs_2019.pdf", ("balance_sheet",))["balance_sheet"]
    assert balance.unit == "US dollars"
    spec = ModelSpec(company="x", unit="millions of Rupiah", statements={})
    with pytest.raises(ValueError, match="currency"):
        spec.scale_to_target(balance.unit)
    assert spec.scale_to_target("thousands") == Decimal("0.001")


def test_a_notes_page_titled_with_a_cash_flow_note_is_not_the_cash_flow():
    """DEWA 2021 note 34, "Cash Flow Additional Information", sits in the
    header band of a notes page, and its debt table was read as cash flow."""
    from ir3s.extractor import find_statement_pages

    path = AUDIT / "dewa" / "dewa_fs_2021.pdf"
    if not path.exists():
        pytest.skip(f"{path.name} not present")
    assert find_statement_pages(pymupdf.open(str(path)))["cash_flow"] == [10]


def test_a_negative_printed_with_spaced_parentheses_is_negative():
    """DEWA 2025 prints "(  899.579.380 )": the "(" is a word of its own.

    The figure read as positive, so every expense and payment flipped sign.
    """
    statements = _statements(AUDIT / "dewa" / "dewa_fs_2025.pdf")
    income = statements["income_statement"]
    assert _row_with(income, -5430580843).value_for(2024) == -5592491287
    cash_flow = statements["cash_flow"]
    employees = _row_with(cash_flow, -1047532373)
    assert employees.value_for(2024) == -899579380
    investing = _row_with(cash_flow, -1686223847)
    assert investing is not None


def test_a_negative_whose_closing_parenthesis_is_a_word_of_its_own_is_kept():
    """CASS 2024 prints the 2023 column as "(1.650.665" then ")" apart.

    The figure did not read as a number, so the 2023 operating expenses,
    other operating expenses and income tax were dropped.
    """
    income = _statements(AUDIT / "cass" / "cass_fs_2024.pdf", ("income_statement",))["income_statement"]
    assert _row_with(income, -1947427).value_for(2023) == -1650665
    assert _row_with(income, -203130).value_for(2023) == -93101
    assert _row_with(income, -188955).value_for(2023) == -135919
    assert ")" not in _row_with(income, -1947427).label


def test_a_short_figure_beside_the_notes_column_is_a_figure():
    """CASS 2024 prints "61" for assets held for sale, right-aligned in 2024.

    Its left edge falls near the "Notes" header, so it was read as a note
    reference and the current assets no longer added up to their total.
    """
    balance = _statements(AUDIT / "cass" / "cass_fs_2024.pdf", ("balance_sheet",))["balance_sheet"]
    held = next(row for row in balance.rows if "held for sale" in row.label and row.value_for(2023) == 230)
    assert held.value_for(2024) == 61
    assert held.notes == "9"


def test_the_parent_entity_statements_at_the_back_are_not_the_consolidated_ones():
    """PBSA 2024 appends "(ENTITAS INDUK SAJA)" statements from page 117.

    They carry a Notes column too, so they were read as more pages of the
    consolidated statements and their lines were mixed in.
    """
    statements = _statements(AUDIT / "prda" / "pbsa_fs_2024.pdf")
    for statement in statements.values():
        assert max(row.page for row in statement.rows) < 100
